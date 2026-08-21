"""Extreme Weather Events Report routes"""
from flask import Blueprint, render_template, request, send_file, flash, redirect, url_for, jsonify
from flask_login import login_required, current_user
from app.models import Station, Parameter, Observation, WordDocExtremeData
from app import db
from datetime import datetime
import io
import os

# For Word export
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# For PDF export
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER

extreme_weather_bp = Blueprint('extreme_weather', __name__)

def get_month_name(month_num):
    months = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September', 'October', 'November', 'December']
    if 1 <= month_num <= 12:
        return months[month_num - 1].upper()
    return ''

def format_val_date(obs):
    if not obs or obs.obs_value is None:
        return "-"
    val_str = f"{float(obs.obs_value):.1f}"
    if obs.date_of_extreme and obs.date_of_extreme.strip() not in ['-', ',', 'None', '']:
        return f"{val_str} ({obs.date_of_extreme.strip()})"
    return val_str

def get_reports_for_month(month, start_year=None, end_year=None):
    """Fetch observations and group by station and year for the given month"""
    stations = Station.query.filter_by(is_active=True).order_by(Station.station_name).all()
    
    p_tmax = Parameter.query.filter_by(parameter_code='TEMP_MAX_HIGH').first()
    p_tmin = Parameter.query.filter_by(parameter_code='TEMP_MIN_LOW').first()
    p_rf24 = Parameter.query.filter_by(parameter_code='RAINFALL_MAX_24HR').first()
    p_rfm = Parameter.query.filter_by(parameter_code='RAINFALL_TOTAL').first()
    p_ids = [p.parameter_id for p in [p_tmax, p_tmin, p_rf24, p_rfm] if p]

    reports_by_station = []
    
    for station in stations:
        query = Observation.query.filter(
            Observation.station_id == station.station_id,
            Observation.obs_month == month,
            Observation.parameter_id.in_(p_ids)
        )
        
        if start_year:
            query = query.filter(Observation.obs_year >= start_year)
        if end_year:
            query = query.filter(Observation.obs_year <= end_year)
            
        observations = query.order_by(Observation.obs_year.asc()).all()
        
        # If no observations and no all-time record, maybe still show empty table,
        # but to keep it clean, we show all active stations.
        
        # WordDoc Data
        doc_query = WordDocExtremeData.query.filter_by(station_id=station.station_id, obs_month=month)
        if start_year: doc_query = doc_query.filter(WordDocExtremeData.obs_year >= start_year)
        if end_year: doc_query = doc_query.filter(WordDocExtremeData.obs_year <= end_year)
        doc_observations = doc_query.order_by(WordDocExtremeData.obs_year.asc()).all()

        years_data = {}
        
        # 1. Populate with Master Data (Baseline)
        for obs in observations:
            y = obs.obs_year
            if y not in years_data:
                years_data[y] = {'year': y, 'tmax_obs': None, 'tmin_obs': None, 'rf24_obs': None, 'rfm_obs': None}
            if p_tmax and obs.parameter_id == p_tmax.parameter_id: years_data[y]['tmax_obs'] = obs
            elif p_tmin and obs.parameter_id == p_tmin.parameter_id: years_data[y]['tmin_obs'] = obs
            elif p_rf24 and obs.parameter_id == p_rf24.parameter_id: years_data[y]['rf24_obs'] = obs
            elif p_rfm and obs.parameter_id == p_rfm.parameter_id: years_data[y]['rfm_obs'] = obs
            
        # 2. Overwrite with Word Doc Data (Priority - contains accurate dates for 10 years)
        for obs in doc_observations:
            y = obs.obs_year
            if y not in years_data:
                years_data[y] = {'year': y, 'tmax_obs': None, 'tmin_obs': None, 'rf24_obs': None, 'rfm_obs': None}
            if p_tmax and obs.parameter_id == p_tmax.parameter_id: years_data[y]['tmax_obs'] = obs
            elif p_tmin and obs.parameter_id == p_tmin.parameter_id: years_data[y]['tmin_obs'] = obs
            elif p_rf24 and obs.parameter_id == p_rf24.parameter_id: years_data[y]['rf24_obs'] = obs
            elif p_rfm and obs.parameter_id == p_rfm.parameter_id: years_data[y]['rfm_obs'] = obs
                
        report_data = [years_data[y] for y in sorted(years_data.keys(), reverse=True)]
        
        # Master ATR
        atr_tmax_m = Observation.query.filter_by(station_id=station.station_id, obs_month=month, parameter_id=p_tmax.parameter_id).filter(Observation.obs_value != None).order_by(Observation.obs_value.desc()).first() if p_tmax else None
        atr_tmin_m = Observation.query.filter_by(station_id=station.station_id, obs_month=month, parameter_id=p_tmin.parameter_id).filter(Observation.obs_value != None).order_by(Observation.obs_value.asc()).first() if p_tmin else None
        atr_rf24_m = Observation.query.filter_by(station_id=station.station_id, obs_month=month, parameter_id=p_rf24.parameter_id).filter(Observation.obs_value != None).order_by(Observation.obs_value.desc()).first() if p_rf24 else None
        atr_rfm_m = Observation.query.filter_by(station_id=station.station_id, obs_month=month, parameter_id=p_rfm.parameter_id).filter(Observation.obs_value != None).order_by(Observation.obs_value.desc()).first() if p_rfm else None
        
        # Doc ATR
        atr_tmax_d = WordDocExtremeData.query.filter_by(station_id=station.station_id, obs_month=month, parameter_id=p_tmax.parameter_id).filter(WordDocExtremeData.obs_value != None).order_by(WordDocExtremeData.obs_value.desc()).first() if p_tmax else None
        atr_tmin_d = WordDocExtremeData.query.filter_by(station_id=station.station_id, obs_month=month, parameter_id=p_tmin.parameter_id).filter(WordDocExtremeData.obs_value != None).order_by(WordDocExtremeData.obs_value.asc()).first() if p_tmin else None
        atr_rf24_d = WordDocExtremeData.query.filter_by(station_id=station.station_id, obs_month=month, parameter_id=p_rf24.parameter_id).filter(WordDocExtremeData.obs_value != None).order_by(WordDocExtremeData.obs_value.desc()).first() if p_rf24 else None
        atr_rfm_d = WordDocExtremeData.query.filter_by(station_id=station.station_id, obs_month=month, parameter_id=p_rfm.parameter_id).filter(WordDocExtremeData.obs_value != None).order_by(WordDocExtremeData.obs_value.desc()).first() if p_rfm else None
        
        def get_max_atr(m, d):
            if m and d: return m if m.obs_value > d.obs_value else d
            return m or d

        def get_min_atr(m, d):
            if m and d: return m if m.obs_value < d.obs_value else d
            return m or d

        all_time_record = {
            'tmax': get_max_atr(atr_tmax_m, atr_tmax_d),
            'tmin': get_min_atr(atr_tmin_m, atr_tmin_d),
            'rf24': get_max_atr(atr_rf24_m, atr_rf24_d),
            'rfm': get_max_atr(atr_rfm_m, atr_rfm_d)
        }
        
        reports_by_station.append({
            'station': station,
            'report_data': report_data,
            'all_time_record': all_time_record
        })
        
    return reports_by_station

@extreme_weather_bp.route('/', methods=['GET', 'POST'])
@login_required
def index():
    """Extreme Weather Events Report"""
    selected_month = None
    start_year = None
    end_year = None
    reports_by_station = []
    
    if request.method == 'POST':
        selected_month = request.form.get('month')
        start_year = request.form.get('start_year')
        end_year = request.form.get('end_year')
        
        if selected_month:
            selected_month = int(selected_month)
            sy = int(start_year) if start_year else None
            ey = int(end_year) if end_year else None
            
            # Pass back the exact values for the form
            start_year = sy
            end_year = ey
            
            reports_by_station = get_reports_for_month(selected_month, sy, ey)

    return render_template(
        'reports/extreme_weather.html', 
        selected_month=selected_month,
        start_year=start_year,
        end_year=end_year,
        reports_by_station=reports_by_station
    )

@extreme_weather_bp.route('/api/update_extreme_date', methods=['POST'])
@login_required
def update_extreme_date():
    """AJAX endpoint to update date_of_extreme"""
    if not current_user.has_permission('write'):
        return jsonify({'success': False, 'message': 'Permission denied'})
        
    obs_id = request.form.get('observation_id')
    date_val = request.form.get('date_val')
    obs_type = request.form.get('obs_type')
    
    if not obs_id or not obs_type:
        return jsonify({'success': False, 'message': 'Observation ID and type required'})
        
    if obs_type == 'Observation':
        obs = Observation.query.get(obs_id)
    elif obs_type == 'WordDocExtremeData':
        obs = WordDocExtremeData.query.get(obs_id)
    else:
        return jsonify({'success': False, 'message': 'Invalid observation type'})
        
    if not obs:
        return jsonify({'success': False, 'message': 'Observation not found'})
        
    try:
        if date_val:
            obs.date_of_extreme = str(date_val).strip()
        else:
            obs.date_of_extreme = None
            
        obs.updated_at = datetime.utcnow()
        if hasattr(obs, 'updated_by'):
            obs.updated_by = current_user.user_id
        db.session.commit()
        return jsonify({'success': True, 'date_val': obs.date_of_extreme})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)})

@extreme_weather_bp.route('/api/update_extreme_dates_batch', methods=['POST'])
@login_required
def update_extreme_dates_batch():
    """AJAX endpoint to update multiple date_of_extreme records"""
    if not current_user.has_permission('write'):
        return jsonify({'success': False, 'message': 'Permission denied'})
        
    data = request.get_json()
    if not data or not isinstance(data, list):
        return jsonify({'success': False, 'message': 'Invalid data format'})
        
    try:
        for item in data:
            obs_id = item.get('observation_id')
            date_val = item.get('date_val')
            obs_type = item.get('obs_type')
            
            if not obs_id or not obs_type:
                continue
                
            if obs_type == 'Observation':
                obs = Observation.query.get(obs_id)
            elif obs_type == 'WordDocExtremeData':
                obs = WordDocExtremeData.query.get(obs_id)
            else:
                continue
                
            if obs:
                if date_val:
                    obs.date_of_extreme = str(date_val).strip()
                else:
                    obs.date_of_extreme = None
                    
                obs.updated_at = datetime.utcnow()
                if hasattr(obs, 'updated_by'):
                    obs.updated_by = current_user.user_id
                    
        db.session.commit()
        return jsonify({'success': True})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': str(e)})


@extreme_weather_bp.route('/export/word', methods=['POST'])
@login_required
def export_word():
    month = request.form.get('month')
    start_year = request.form.get('start_year')
    end_year = request.form.get('end_year')
    
    if not month:
        flash('Invalid export parameters.', 'danger')
        return redirect(url_for('extreme_weather.index'))
        
    month = int(month)
    sy = int(start_year) if start_year else None
    ey = int(end_year) if end_year else None
    month_name = get_month_name(month)
    
    reports_by_station = get_reports_for_month(month, sy, ey)
    
    doc = Document()
    
    for idx, data in enumerate(reports_by_station):
        station = data['station']
        report_data = data['report_data']
        all_time_record = data['all_time_record']
        
        if idx > 0:
            doc.add_page_break()
            
        # Title formatting
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title.add_run(f"{idx + 1}. {station.station_name.upper()}")
        title_run.bold = True
        title_run.underline = True
        title_run.font.size = Pt(14)
        
        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = subtitle.add_run(f"EXTREME WEATHER EVENTS IN THE MONTH OF {month_name}")
        sub_run.bold = True
        
        doc.add_paragraph() # Spacer
        
        # Table
        table = doc.add_table(rows=2, cols=5)
        table.style = 'Table Grid'
        
        # Merge headers
        c00 = table.cell(0, 0)
        c00.merge(table.cell(1, 0))
        c00.text = "Year"
        
        c01 = table.cell(0, 1)
        c01.merge(table.cell(0, 2))
        c01.text = "Temperature(°C)"
        
        c03 = table.cell(0, 3)
        c03.merge(table.cell(0, 4))
        c03.text = "Rainfall (mm)"
        
        table.cell(1, 1).text = "Highest Maximum(Date)"
        table.cell(1, 2).text = "Lowest Minimum(Date)"
        table.cell(1, 3).text = "24 Hours Highest (Date)"
        table.cell(1, 4).text = "Monthly Total"
        
        # Format headers
        for r in range(2):
            for c in range(5):
                cell = table.cell(r, c)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.bold = True
            
        for row_data in report_data:
            row_cells = table.add_row().cells
            row_cells[0].text = str(row_data['year'])
            row_cells[1].text = format_val_date(row_data['tmax_obs'])
            row_cells[2].text = format_val_date(row_data['tmin_obs'])
            row_cells[3].text = format_val_date(row_data['rf24_obs'])
            row_cells[4].text = f"{float(row_data['rfm_obs'].obs_value):.1f}" if row_data['rfm_obs'] and row_data['rfm_obs'].obs_value is not None else "-"
            
            for cell in row_cells:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                
        # All Time Records as the last row
        if any(all_time_record.values()):
            atr_cells = table.add_row().cells
            atr_cells[0].text = "ALL TIME RECORD"
            
            def fmt_atr(obs):
                if not obs or obs.obs_value is None: return "-"
                v_str = f"{float(obs.obs_value):.1f}"
                if obs.date_of_extreme or obs.obs_year:
                    d_clean = obs.date_of_extreme.strip() if obs.date_of_extreme and obs.date_of_extreme.strip() not in ['-', ',', 'None', ''] else None
                    if d_clean and obs.obs_year:
                        d_str = f"({d_clean}, {obs.obs_year})"
                    elif d_clean:
                        d_str = f"({d_clean})"
                    elif obs.obs_year:
                        d_str = f"({obs.obs_year})"
                    else:
                        d_str = ""
                    return f"{v_str} {d_str}".strip()
                return v_str
                
            atr_cells[1].text = fmt_atr(all_time_record['tmax'])
            atr_cells[2].text = fmt_atr(all_time_record['tmin'])
            atr_cells[3].text = fmt_atr(all_time_record['rf24'])
            
            obs_rfm = all_time_record['rfm']
            if obs_rfm and obs_rfm.obs_value is not None:
                atr_cells[4].text = f"{float(obs_rfm.obs_value):.1f} ({obs_rfm.obs_year})" if obs_rfm.obs_year else f"{float(obs_rfm.obs_value):.1f}"
            else:
                atr_cells[4].text = "-"
                
            for cell in atr_cells:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True

    import tempfile
    import os

    # Create temporary path for modern .docx format (consistent with Import Extreme and no COM automation needed)
    temp_dir = tempfile.gettempdir()
    docx_path = os.path.join(temp_dir, f'temp_export_{month_name}.docx')
    
    # Save as .docx
    doc.save(docx_path)
    
    with open(docx_path, 'rb') as f:
        output_bytes = f.read()
        
    # Clean up temp file
    try:
        if os.path.exists(docx_path):
            os.remove(docx_path)
    except:
        pass
        
    output = io.BytesIO(output_bytes)
    output.seek(0)
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        as_attachment=True,
        download_name=f'Extreme_Weather_All_Stations_{month_name}.docx'
    )

@extreme_weather_bp.route('/export/pdf', methods=['POST'])
@login_required
def export_pdf():
    month = request.form.get('month')
    start_year = request.form.get('start_year')
    end_year = request.form.get('end_year')
    
    if not month:
        flash('Invalid export parameters.', 'danger')
        return redirect(url_for('extreme_weather.index'))
        
    month = int(month)
    sy = int(start_year) if start_year else None
    ey = int(end_year) if end_year else None
    month_name = get_month_name(month)
    
    reports_by_station = get_reports_for_month(month, sy, ey)
    
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    elements = []
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('Title', parent=styles['Heading1'], alignment=TA_CENTER)
    subtitle_style = ParagraphStyle('Subtitle', parent=styles['Heading2'], alignment=TA_CENTER)
    
    for idx, data in enumerate(reports_by_station):
        station = data['station']
        report_data = data['report_data']
        all_time_record = data['all_time_record']
        
        if idx > 0:
            elements.append(PageBreak())
            
        title_para = Paragraph(f"<u>{idx + 1}. {station.station_name.upper()}</u>", title_style)
        elements.append(title_para)
        elements.append(Paragraph(f"EXTREME WEATHER EVENTS IN THE MONTH OF {month_name}", subtitle_style))
        elements.append(Spacer(1, 20))
        
        # Table data
        table_data = [
            ['Year', 'Temperature(°C)', '', 'Rainfall (mm)', ''],
            ['', 'Highest\nMaximum(Date)', 'Lowest\nMinimum(Date)', '24 Hours Highest\n(Date)', 'Monthly Total']
        ]
        
        for row_data in report_data:
            rfm_val = f"{float(row_data['rfm_obs'].obs_value):.1f}" if row_data['rfm_obs'] and row_data['rfm_obs'].obs_value is not None else "-"
            table_data.append([
                str(row_data['year']),
                format_val_date(row_data['tmax_obs']),
                format_val_date(row_data['tmin_obs']),
                format_val_date(row_data['rf24_obs']),
                rfm_val
            ])
            
        if any(all_time_record.values()):
            def fmt_atr_pdf(obs):
                if not obs or obs.obs_value is None: return "-"
                v_str = f"{float(obs.obs_value):.1f}"
                if obs.date_of_extreme or obs.obs_year:
                    d_clean = obs.date_of_extreme.strip() if obs.date_of_extreme and obs.date_of_extreme.strip() not in ['-', ',', 'None', ''] else None
                    if d_clean and obs.obs_year:
                        d_str = f"({d_clean}, {obs.obs_year})"
                    elif d_clean:
                        d_str = f"({d_clean})"
                    elif obs.obs_year:
                        d_str = f"({obs.obs_year})"
                    else:
                        d_str = ""
                    return f"{v_str} {d_str}".strip()
                return v_str
                
            obs_mrf = all_time_record['rfm']
            mrf_str = "-"
            if obs_mrf and obs_mrf.obs_value is not None:
                mrf_str = f"{float(obs_mrf.obs_value):.1f} ({obs_mrf.obs_year})" if obs_mrf.obs_year else f"{float(obs_mrf.obs_value):.1f}"
            
            table_data.append([
                "ALL TIME\nRECORD",
                fmt_atr_pdf(all_time_record['tmax']),
                fmt_atr_pdf(all_time_record['tmin']),
                fmt_atr_pdf(all_time_record['rf24']),
                mrf_str
            ])
            
        t = Table(table_data, repeatRows=2)
        
        style_cmds = [
            ('SPAN', (0, 0), (0, 1)),
            ('SPAN', (1, 0), (2, 0)),
            ('SPAN', (3, 0), (4, 0)),
            ('BACKGROUND', (0,0), (-1,1), colors.lightgrey),
            ('TEXTCOLOR', (0,0), (-1,1), colors.black),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('FONTNAME', (0,0), (-1,1), 'Helvetica-Bold'),
            ('GRID', (0,0), (-1,-1), 1, colors.black)
        ]
        
        if any(all_time_record.values()):
            style_cmds.append(('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'))
            
        t.setStyle(TableStyle(style_cmds))
        
        elements.append(t)
        elements.append(Spacer(1, 30))
            
    doc.build(elements)
    output.seek(0)
    
    return send_file(
        output,
        mimetype='application/pdf',
        as_attachment=True,
        download_name=f'Extreme_Weather_All_Stations_{month_name}.pdf'
    )
