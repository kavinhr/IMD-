"""Reports routes"""
from flask import Blueprint, render_template, request, send_file, flash, redirect, url_for
from flask_login import login_required
from app.models import Station, Parameter, Observation
import io, csv
from datetime import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

reports_bp = Blueprint('reports', __name__)

@reports_bp.route('/')
@login_required
def index():
    """Reports interface"""
    stations = Station.query.filter_by(is_active=True).order_by(Station.station_name).all()
    parameters = Parameter.query.filter_by(is_active=True).order_by(Parameter.display_order).all()
    return render_template('reports/index.html', stations=stations, parameters=parameters)

@reports_bp.route('/export/csv', methods=['POST'])
@login_required
def export_csv():
    """Export data to CSV"""
    station_ids = request.form.getlist('station_id')
    parameter_ids = request.form.getlist('parameter_id')
    start_year = request.form.get('start_year')
    end_year = request.form.get('end_year')
    
    query = Observation.query.join(Station).join(Parameter)
    
    station_ids = [s for s in station_ids if s]
    if station_ids:
        query = query.filter(Observation.station_id.in_(station_ids))
        
    parameter_ids = [p for p in parameter_ids if p]
    if parameter_ids:
        query = query.filter(Observation.parameter_id.in_(parameter_ids))
        
    if start_year:
        query = query.filter(Observation.obs_year >= int(start_year))
    if end_year:
        query = query.filter(Observation.obs_year <= int(end_year))
    
    observations = query.all()
    
    # Create CSV
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['Station', 'Parameter', 'Year', 'Month', 'Value', 'Unit', 'Quality'])
    
    for obs in observations:
        val = obs.obs_value
        if val is not None:
            val_str = f"{float(val):.1f}"
        else:
            val_str = obs.value_text if obs.value_text else ''
            
        writer.writerow([
            obs.station.station_name,
            obs.parameter.parameter_name,
            obs.obs_year,
            obs.obs_month,
            val_str,
            obs.parameter.unit,
            obs.data_quality_flag
        ])
    
    output.seek(0)
    return send_file(
        io.BytesIO(output.getvalue().encode()),
        mimetype='text/csv',
        as_attachment=True,
        download_name=f'weather_data_{datetime.now().strftime("%Y%m%d")}.csv'
    )

@reports_bp.route('/export/excel', methods=['POST'])
@login_required
def export_excel():
    """Export data to Excel in IMD block format"""
    station_ids = request.form.getlist('station_id')
    parameter_ids = request.form.getlist('parameter_id')
    start_year = request.form.get('start_year')
    end_year = request.form.get('end_year')
    
    query = Observation.query.join(Station).join(Parameter)
    
    station_ids = [s for s in station_ids if s]
    if station_ids:
        query = query.filter(Observation.station_id.in_(station_ids))
        
    parameter_ids = [p for p in parameter_ids if p]
    if parameter_ids:
        query = query.filter(Observation.parameter_id.in_(parameter_ids))
        
    if start_year:
        query = query.filter(Observation.obs_year >= int(start_year))
    if end_year:
        query = query.filter(Observation.obs_year <= int(end_year))
    
    # Order by Station, Parameter display_order, Year, Month
    query = query.order_by(
        Station.station_name, 
        Parameter.display_order, 
        Parameter.parameter_id,
        Observation.obs_year, 
        Observation.obs_month
    )
    observations = query.all()
    
    if not observations:
        flash('No data found for the given criteria.', 'warning')
        return redirect(url_for('reports.index'))
        
    # Group data by Station -> Parameter -> Year -> Month
    data_dict = {}
    for obs in observations:
        st_name = obs.station.station_name
        param_obj = obs.parameter
        # Using parameter key to group data accurately
        param_key = (param_obj.parameter_name, param_obj.unit, param_obj.parameter_id)
        
        y = obs.obs_year
        m = obs.obs_month
        
        val = obs.obs_value
        if val is not None:
            val = round(float(val), 1)
        else:
            val = obs.value_text if obs.value_text else ''
            
        if st_name not in data_dict:
            data_dict[st_name] = {}
        if param_key not in data_dict[st_name]:
            data_dict[st_name][param_key] = {}
        if y not in data_dict[st_name][param_key]:
            data_dict[st_name][param_key][y] = {i: '' for i in range(1, 13)}
            
        data_dict[st_name][param_key][y][m] = val
        
    wb = Workbook()
    
    # Remove default sheet
    if len(wb.sheetnames) > 0:
        wb.remove(wb.active)
    
    font_bold = Font(bold=True)
    
    for st_name, st_data in data_dict.items():
        # Sheet name limited to 31 chars and valid chars
        safe_sheet_name = "".join([c for c in st_name if c.isalnum() or c in (' ', '_', '-')])[:31]
        if not safe_sheet_name:
            safe_sheet_name = "Station"
        ws = wb.create_sheet(title=safe_sheet_name)
        
        # Write Station Name at the top
        st_cell = ws.cell(row=2, column=1, value=f"STATION: {st_name.upper()}")
        st_cell.font = Font(bold=True, size=12)
        
        # Initial offset to mirror top empty rows
        current_row = 7
        
        param_index = 1
        for param_key, years_data in st_data.items():
            param_name, unit, _ = param_key
            unit_str = f" ({unit})" if unit else ""
            title = f"{param_index}.  ELEMENT: {param_name}{unit_str}"
            
            # Write Title
            cell = ws.cell(row=current_row, column=1, value=title.upper())
            cell.font = font_bold
            current_row += 2
            
            # Write Header
            headers = ['YEAR', 'JAN', 'FEB', 'MAR', 'APR', 'MAY', 'JUN', 'JUL', 'AUG', 'SEP', 'OCT', 'NOV', 'DEC']
            for col_idx, h in enumerate(headers, start=1):
                c = ws.cell(row=current_row, column=col_idx, value=h)
                c.font = font_bold
            current_row += 1
            
            # Write Data
            for y in sorted(years_data.keys()):
                ws.cell(row=current_row, column=1, value=y)
                monthly_vals = years_data[y]
                for m_idx in range(1, 13):
                    ws.cell(row=current_row, column=m_idx + 1, value=monthly_vals.get(m_idx, ''))
                current_row += 1
                
            current_row += 3 # Empty rows between parameters
            param_index += 1
            
        # Adjust column widths
        ws.column_dimensions['A'].width = 10
        for col_idx in range(2, 14):
            # Convert column index to letter (B=2, C=3, ..., M=13)
            col_letter = chr(64 + col_idx)
            ws.column_dimensions[col_letter].width = 8
            
    # Save to BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'imd_weather_data_{datetime.now().strftime("%Y%m%d")}.xlsx'
    )
